"""Download all business data as one Excel workbook (Settings > Export data)."""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session, selectinload

from ..database import get_db
from ..models import Customer, Invoice, Product
from ..routers.customers import _with_stats
from ..services import invoice_service, settings_service

router = APIRouter(prefix="/api/export", tags=["export"])

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1B3A6B")
MONEY = "#,##0.00"
DATE_FORMAT = "DD-MM-YYYY"

# Business settings worth exporting (logo files and internal counters are not).
SETTINGS_FIELDS = (
    ("business_name", "Business name"), ("address", "Address"), ("phone", "Phone"), ("email", "E-Mail"),
    ("gstin", "GSTIN/UIN"), ("pan", "PAN"), ("state", "State"), ("state_code", "State code"),
    ("bank_name", "Bank name"), ("account_name", "Account holder"), ("account_number", "Account number"),
    ("ifsc", "IFSC"), ("bank_branch", "Branch"), ("upi_id", "UPI ID"), ("invoice_prefix", "Invoice prefix"), ("invoice_suffix", "Invoice suffix"),
    ("invoice_digits", "Digits in invoice number"), ("default_payment_terms_days", "Default payment terms (days)"),
    ("invoice_terms", "Terms & conditions"), ("invoice_notes", "Default notes"),
    ("footer_phone", "Footer: call number"), ("footer_whatsapp", "Footer: WhatsApp number"),
    ("footer_email", "Footer: e-mail"), ("footer_address", "Footer: address"),
)


def _cell(value):
    """Convert a value to something Excel stores faithfully (numbers stay numbers, dates stay dates)."""
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, (date, int, float, bool)) or value is None:
        return value
    return str(value)


def _sheet(wb: Workbook, title: str, headers: list[str], rows, money_cols=(), date_cols=(), first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for cell in ws[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    widths = [len(h) + 2 for h in headers]
    for row in rows:
        ws.append([_cell(v) for v in row])
        for index, value in enumerate(row):
            if isinstance(value, str) and value.startswith("="):
                ws.cell(row=ws.max_row, column=index + 1).data_type = "s"  # text, never evaluated as a formula
            widths[index] = min(max(widths[index], len(str(value if value is not None else "")) + 2), 60)
    for index, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(index)].width = widths[index - 1]
        number_format = MONEY if header in money_cols else DATE_FORMAT if header in date_cols else None
        if number_format:
            for (cell,) in ws.iter_rows(min_row=2, min_col=index, max_col=index):
                cell.number_format = number_format
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


@router.get("/xlsx")
def export_xlsx(db: Session = Depends(get_db)) -> StreamingResponse:
    invoice_service.refresh_overdue(db)
    wb = Workbook()

    settings = settings_service.get_settings(db)
    _sheet(wb, "Business", ["Setting", "Value"], [(label, getattr(settings, key, "")) for key, label in SETTINGS_FIELDS], first=True)
    wb["Business"].column_dimensions["B"].width = 60

    customers = db.scalars(select(Customer).order_by(func.lower(Customer.name))).all()
    _sheet(wb, "Customers", [
        "Name", "Company", "Phone", "E-Mail", "GSTIN", "State", "State code", "Billing address", "Shipping address", "Notes",
        "Invoices", "Total billed", "Total paid", "Outstanding",
    ], [
        (c.name, c.company_name, c.phone, c.email, c.gstin, c.state, c.state_code, c.billing_address, c.shipping_address,
         c.notes, c.invoice_count, c.total_billed, c.total_paid, c.outstanding)
        for c in _with_stats(db, customers)
    ], money_cols=("Total billed", "Total paid", "Outstanding"))

    products = db.scalars(select(Product).order_by(func.lower(Product.name))).all()
    _sheet(wb, "Products", ["Part No", "Description of Goods", "Notes", "HSN/SAC", "Unit", "Price (excl. GST)", "GST %"], [
        (p.sku, p.name, p.description, p.hsn_sac, p.unit, p.price, p.gst_rate) for p in products
    ], money_cols=("Price (excl. GST)",))

    invoices = db.scalars(
        select(Invoice).options(selectinload(Invoice.items)).order_by(Invoice.invoice_date, Invoice.id)
    ).all()
    outs = [invoice_service.to_out(i, db) for i in invoices]
    _sheet(wb, "Invoices", [
        "Invoice No", "Date", "Due date", "Customer", "Customer GSTIN", "Place of supply", "Status", "Subtotal", "Discount",
        "Taxable value", "CGST", "SGST", "IGST", "Total tax", "Round off", "Grand total", "Paid", "Balance due",
        "Other reference", "Phone", "ESN", "Payment terms", "Notes",
    ], [
        (o.invoice_number, o.invoice_date, o.due_date, o.customer.get("display_name", ""), o.customer.get("gstin", ""),
         o.place_of_supply_label or o.place_of_supply, o.status, o.subtotal, o.discount, o.taxable_amount, o.cgst, o.sgst,
         o.igst, o.total_tax, o.round_off, o.grand_total, o.amount_paid, o.balance_due, o.other_reference, o.contact_phone,
         o.esn, o.payment_terms, o.notes)
        for o in outs
    ], money_cols=("Subtotal", "Discount", "Taxable value", "CGST", "SGST", "IGST", "Total tax", "Round off", "Grand total",
                   "Paid", "Balance due"), date_cols=("Date", "Due date"))

    _sheet(wb, "Invoice Items", [
        "Invoice No", "S No", "Part No", "Description of Goods", "Notes", "HSN/SAC", "Quantity", "Unit", "Rate", "Discount",
        "GST %", "Taxable value", "CGST", "SGST", "IGST", "Amount",
    ], [
        (o.invoice_number, n, i.part_no, i.name, i.description, i.hsn_sac, i.quantity, i.unit, i.rate, i.discount, i.gst_rate,
         i.taxable_amount, i.cgst, i.sgst, i.igst, i.amount)
        for o in outs for n, i in enumerate(o.items, start=1)
    ], money_cols=("Rate", "Discount", "Taxable value", "CGST", "SGST", "IGST", "Amount"))

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    name = f"OHM-Powerin-data-{date.today().isoformat()}.xlsx"
    return StreamingResponse(buffer, media_type=XLSX, headers={"Content-Disposition": f'attachment; filename="{name}"'})
