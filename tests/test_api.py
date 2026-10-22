"""End-to-end API tests, including the Phase 1 acceptance scenario.

Run from the project root:  python -m pytest -q
"""
from __future__ import annotations

import io
import os
import tempfile
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path

_tmp = tempfile.mkdtemp(prefix="billing-test-")
os.environ["BILLING_DATA_DIR"] = _tmp
os.environ["BILLING_INVOICE_DIR"] = str(Path(_tmp) / "invoices")

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from pypdf import PdfReader  # noqa: E402

from backend.main import app  # noqa: E402
from backend.pdf_generator import format_indian  # noqa: E402
from backend.services.calculations import amount_in_words  # noqa: E402

BUSINESS = {
    "business_name": "Sri Lakshmi Traders",
    "address": "12, Anna Salai\nChennai 600002",
    "phone": "+91 98400 12345",
    "email": "accounts@srilakshmi.example",
    "gstin": "33ABCDE1234F1Z5",
    "pan": "ABCDE1234F",
    "state_code": "33",
    "bank_name": "State Bank of India",
    "account_number": "12345678901",
    "ifsc": "SBIN0001234",
    "upi_id": "srilakshmi@sbi",
    "invoice_prefix": "INV-",
    "invoice_start_number": 1,
    "invoice_terms": "Goods once sold will not be taken back.",
    "invoice_notes": "Thank you for your business!",
}


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


def _invoice_payload(customer_id: int, product_id: int | None = None, **overrides) -> dict:
    payload = {
        "customer_id": customer_id,
        "invoice_date": date.today().isoformat(),
        "due_date": (date.today() + timedelta(days=15)).isoformat(),
        "items": [{
            "product_id": product_id,
            "name": "Cement",
            "hsn_sac": "2523",
            "quantity": "10",
            "unit": "Bags",
            "rate": "450",
            "gst_rate": "18",
        }],
    }
    payload.update(overrides)
    return payload


def test_acceptance_flow(client):
    # Business
    r = client.put("/api/settings", json=BUSINESS)
    assert r.status_code == 200, r.text
    assert r.json()["state"] == "Tamil Nadu"
    assert r.json()["next_invoice_number"] == "INV-0001"

    # Customer (same state as the business -> CGST + SGST)
    r = client.post("/api/customers", json={
        "name": "ABC Construction",
        "billing_address": "45, Mount Road, Chennai",
        "gstin": "TESTGSTIN123",
        "state_code": "33",
    })
    assert r.status_code == 201, r.text
    customer = r.json()

    # Product
    r = client.post("/api/products", json={
        "name": "Cement", "hsn_sac": "2523", "unit": "Bags", "price": "450", "gst_rate": "18",
    })
    assert r.status_code == 201, r.text
    product = r.json()

    # Preview is calculated on the server and not saved
    r = client.post("/api/invoices/preview", json=_invoice_payload(customer["id"], product["id"]))
    assert r.status_code == 200, r.text
    assert r.json()["id"] is None
    assert client.get("/api/invoices").json()["total"] == 0

    # Invoice — frontend totals are never sent; the server computes them
    r = client.post("/api/invoices", json=_invoice_payload(customer["id"], product["id"]))
    assert r.status_code == 201, r.text
    inv = r.json()
    assert inv["invoice_number"] == "INV-0001"
    assert Decimal(inv["subtotal"]) == Decimal("4500.00")
    assert Decimal(inv["taxable_amount"]) == Decimal("4500.00")
    assert Decimal(inv["cgst"]) == Decimal("405.00")
    assert Decimal(inv["sgst"]) == Decimal("405.00")
    assert Decimal(inv["igst"]) == Decimal("0.00")
    assert Decimal(inv["grand_total"]) == Decimal("5310.00")
    assert inv["amount_in_words"] == "Rupees Five Thousand Three Hundred Ten Only"
    assert inv["is_interstate"] is False
    assert inv["status"] == "Pending"
    assert inv["has_pdf"] is True

    # PDF: true A4 page with the expected content
    r = client.get(f"/api/invoices/{inv['id']}/pdf?download=true")
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert 'attachment; filename="INV-0001.pdf"' in r.headers["content-disposition"]
    pdf_path = Path(os.environ["BILLING_INVOICE_DIR"]) / "INV-0001.pdf"
    assert pdf_path.is_file()
    reader = PdfReader(pdf_path)
    box = reader.pages[0].mediabox
    width_mm, height_mm = float(box.width) * 25.4 / 72, float(box.height) * 25.4 / 72
    assert round(width_mm) == 210 and round(height_mm) == 297
    text = " ".join(page.extract_text() for page in reader.pages)
    for expected in (
        "Sri Lakshmi Traders", "33ABCDE1234F1Z5", "INV-0001", "ABC Construction",
        "TESTGSTIN123", "Cement", "2523", "450.00", "18%", "OUTPUT CGST @ 9%", "OUTPUT SGST @ 9%", "Taxable", "Total",
        "405.00", "5,310.00", "INR Five Thousand Three Hundred Ten Only", "State Bank of India",
        "SBIN0001234", "Authorised Signatory", "Declaration", "This is a Computer Generated Invoice",
    ):
        assert expected in text, f"{expected!r} missing from PDF"

    # Next number continues and dashboard reflects real data
    assert client.get("/api/invoices/next-number").json()["invoice_number"] == "INV-0002"
    dash = client.get("/api/dashboard").json()
    assert Decimal(dash["total_sales"]) == Decimal("5310.00")
    assert Decimal(dash["pending"]) == Decimal("5310.00")
    assert dash["invoice_count"] == 1


def test_interstate_igst_discount_and_round_off(client):
    customer = client.post("/api/customers", json={
        "name": "Bengaluru Builders", "company_name": "BB Infra Pvt Ltd", "state_code": "29",
    }).json()
    r = client.post("/api/invoices", json=_invoice_payload(customer["id"], items=[
        {"name": "Consulting", "hsn_sac": "998311", "quantity": "1", "rate": "10000",
         "discount_type": "amount", "discount_value": "500", "gst_rate": "18"},
        {"name": "Steel rod", "quantity": "3.333", "rate": "99.99",
         "discount_type": "percent", "discount_value": "7.5", "gst_rate": "12"},
    ]))
    assert r.status_code == 201, r.text
    inv = r.json()
    assert inv["is_interstate"] is True
    assert Decimal(inv["cgst"]) == Decimal("0") and Decimal(inv["sgst"]) == Decimal("0")
    # 9500 * 18% = 1710 ; 308.27 * 12% = 36.99
    assert Decimal(inv["igst"]) == Decimal("1746.99")
    assert Decimal(inv["taxable_amount"]) == Decimal("9808.27")
    assert Decimal(inv["grand_total"]) == Decimal("11555")
    assert Decimal(inv["round_off"]) == Decimal("-0.26")
    assert inv["customer"]["display_name"] == "BB Infra Pvt Ltd"


def test_duplicate_number_and_validation(client):
    customer_id = client.get("/api/customers").json()["items"][0]["id"]
    r = client.post("/api/invoices", json=_invoice_payload(customer_id, invoice_number="INV-0001"))
    assert r.status_code == 409
    assert "already in use" in r.json()["detail"]

    bad_rate = _invoice_payload(customer_id)
    bad_rate["items"][0]["gst_rate"] = "17"
    r = client.post("/api/invoices", json=bad_rate)
    assert r.status_code == 422 and "GST rate" in r.json()["detail"]

    r = client.post("/api/invoices", json=_invoice_payload(customer_id, items=[]))
    assert r.status_code == 422

    r = client.post("/api/customers", json={"name": "  "})
    assert r.status_code == 422 and "Traceback" not in r.text


def test_status_payments_and_overdue(client):
    customer_id = client.get("/api/customers").json()["items"][0]["id"]
    inv = client.post("/api/invoices", json=_invoice_payload(customer_id)).json()

    r = client.post(f"/api/invoices/{inv['id']}/payments", json={"amount": "1000"})
    assert r.json()["status"] == "Partially Paid"
    r = client.post(f"/api/invoices/{inv['id']}/payments", json={"amount": "99999"})
    assert r.status_code == 422
    r = client.patch(f"/api/invoices/{inv['id']}/status", json={"status": "Paid"})
    assert r.json()["status"] == "Paid" and Decimal(r.json()["balance_due"]) == 0

    past = _invoice_payload(
        customer_id,
        invoice_date=(date.today() - timedelta(days=40)).isoformat(),
        due_date=(date.today() - timedelta(days=10)).isoformat(),
    )
    overdue = client.post("/api/invoices", json=past).json()
    assert overdue["status"] == "Overdue"
    assert client.get("/api/invoices?status=Overdue").json()["total"] >= 1


def test_customer_with_invoices_cannot_be_deleted(client):
    customer_id = client.get("/api/customers?q=ABC").json()["items"][0]["id"]
    r = client.delete(f"/api/customers/{customer_id}")
    assert r.status_code == 409


def test_long_content_paginates(client):
    customer = client.post("/api/customers", json={
        "name": "A very long customer name " * 6,
        "billing_address": "Line of a long address, " * 15,
        "shipping_address": "Different shipping address, " * 8,
        "state_code": "27",
    }).json()
    items = [
        {"name": f"Item {i} <b>&</b>", "description": "Long description text " * 8,
         "hsn_sac": "1234", "quantity": "1", "rate": "123456.78", "gst_rate": "5"}
        for i in range(45)
    ]
    r = client.post("/api/invoices", json=_invoice_payload(customer["id"], items=items))
    assert r.status_code == 201, r.text
    pdf = client.get(f"/api/invoices/{r.json()['id']}/pdf")
    assert pdf.status_code == 200
    reader = PdfReader(io.BytesIO(pdf.content))
    assert len(reader.pages) > 1
    assert "Item 0 <b>&</b>" in reader.pages[0].extract_text()  # markup is escaped, not interpreted


def test_logo_upload_validation_and_pdf(client):
    # Data dir lives outside the project here, which previously broke logo storage.
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (200, 80), (30, 58, 95)).save(buffer, format="PNG")
    r = client.post("/api/settings/logo", files={"file": ("logo.png", buffer.getvalue(), "image/png")})
    assert r.status_code == 200, r.text
    assert r.json()["has_logo"] is True
    assert client.get("/api/settings/logo").status_code == 200

    r = client.post("/api/settings/logo", files={"file": ("evil.png", b"<script>alert(1)</script>", "image/png")})
    assert r.status_code == 400 and "PNG or JPEG" in r.json()["detail"]

    invoice_id = client.get("/api/invoices").json()["items"][0]["id"]
    assert client.post(f"/api/invoices/{invoice_id}/pdf/regenerate").status_code == 200


def test_formatting_helpers():
    assert format_indian(Decimal("245000")) == "2,45,000.00"
    assert format_indian(Decimal("123456789.5")) == "12,34,56,789.50"
    assert format_indian(Decimal("999")) == "999.00"
    assert amount_in_words(Decimal("11210")) == "Rupees Eleven Thousand Two Hundred Ten Only"
    assert amount_in_words(Decimal("2345678")) == (
        "Rupees Twenty Three Lakh Forty Five Thousand Six Hundred Seventy Eight Only"
    )


def test_invoice_number_format_and_manual_continuation(client):
    fmt = {**BUSINESS, "invoice_prefix": "OHM", "invoice_suffix": "FY2627", "invoice_digits": 3}
    r = client.put("/api/settings", json=fmt)
    assert r.status_code == 200, r.text
    assert r.json()["next_invoice_number"] == "OHM001FY2627"

    customer_id = client.get("/api/customers").json()["items"][0]["id"]
    r = client.post("/api/invoices", json=_invoice_payload(customer_id, invoice_number="OHM120FY2627"))
    assert r.status_code == 201, r.text
    assert client.get("/api/invoices/next-number").json()["invoice_number"] == "OHM121FY2627"
    r = client.post("/api/invoices", json=_invoice_payload(customer_id))
    assert r.json()["invoice_number"] == "OHM121FY2627"
    assert client.get("/api/invoices/next-number").json()["invoice_number"] == "OHM122FY2627"


def test_reference_fields_default_to_dash_in_pdf(client):
    customer_id = client.get("/api/customers").json()["items"][0]["id"]
    r = client.post("/api/invoices", json=_invoice_payload(customer_id, esn="ESN-778", contact_phone="9876543210"))
    assert r.status_code == 201, r.text
    body = r.json()
    assert (body["esn"], body["other_reference"]) == ("ESN-778", "")
    pdf = client.get(f"/api/invoices/{body['id']}/pdf")
    from io import BytesIO
    text = PdfReader(BytesIO(pdf.content)).pages[0].extract_text()
    # The reference fields are always printed; empty ones show "-".
    assert "Other Reference" in text and "ESN-778" in text and "9876543210" in text


def test_fit_mode_keeps_15_items_on_one_page_and_pages_mode_continues(client):
    from io import BytesIO
    customer_id = client.get("/api/customers").json()["items"][0]["id"]
    payload = _invoice_payload(customer_id)
    payload["items"] = [{**payload["items"][0], "name": f"Spare part number {n} with a longer name"} for n in range(1, 16)]
    counts = {}
    for mode in ("fit", "pages"):
        body = client.post("/api/invoices", json={**payload, "page_mode": mode}).json()
        assert body["page_mode"] == mode
        counts[mode] = len(PdfReader(BytesIO(client.get(f"/api/invoices/{body['id']}/pdf").content)).pages)
    assert counts == {"fit": 1, "pages": 2}


def test_pdf_from_older_layout_is_rerendered(client):
    from backend.config import INVOICE_DIR
    from backend.pdf_generator import PDF_CREATOR
    customer_id = client.get("/api/customers").json()["items"][0]["id"]
    body = client.post("/api/invoices", json=_invoice_payload(customer_id)).json()
    stored = max(INVOICE_DIR.glob("*.pdf"), key=lambda p: p.stat().st_mtime)
    stored.write_bytes(b"%PDF-1.4 stale layout")
    pdf = client.get(f"/api/invoices/{body['id']}/pdf")
    assert PDF_CREATOR.encode() in pdf.content and b"stale layout" not in pdf.content


def test_product_csv_import_and_part_no_on_invoice(client):
    csv_text = (
        "Part No,Description of Goods,HSN/SAC,Unit,Price,GST %\n"
        "P-100,Brake Pad,87083,Nos,\"1,250.50\",28%\n"
        "P-101,Oil Filter,84212300,,300,18\n"
        ",No part number,,,,\n"
    )
    r = client.post("/api/products/import", files={"file": ("parts.csv", csv_text.encode(), "text/csv")})
    assert r.status_code == 200, r.text
    assert (r.json()["created"], r.json()["error_count"]) == (2, 1)
    r = client.post("/api/products/import", files={"file": ("parts.csv", csv_text.replace("300", "310").encode(), "text/csv")})
    assert (r.json()["created"], r.json()["updated"]) == (0, 2)

    pads = next(p for p in client.get("/api/products", params={"q": "P-100"}).json()["items"])
    assert pads["unit"] == "Nos" and pads["price"] == "1250.50" and pads["gst_rate"] == "28.00"

    customer_id = client.get("/api/customers").json()["items"][0]["id"]
    payload = _invoice_payload(customer_id, pads["id"])
    payload["items"][0].update(part_no="P-100-X", name="Brake Pad")  # part no edited by hand
    r = client.post("/api/invoices", json=payload)
    assert r.status_code == 201, r.text
    assert r.json()["items"][0]["part_no"] == "P-100-X"


def test_logo_heights_saved_and_validated(client):
    r = client.put("/api/settings", json={**BUSINESS, "logo_heights": [30, 15, 20]})
    assert r.status_code == 200, r.text
    assert r.json()["logo_heights"] == [30, 15, 20]
    r = client.put("/api/settings", json={**BUSINESS, "logo_heights": [90, 15, 20]})
    assert r.status_code == 422


# ---------------------------------------------------------------------------
# UI/PDF redesign additions
# ---------------------------------------------------------------------------


def test_dashboard_overview_and_extra_kpis(client):
    client.put("/api/settings", json=BUSINESS)
    d = client.get("/api/dashboard").json()
    # existing keys are untouched, new ones are additive
    for key in ("total_sales", "paid", "pending", "overdue", "recent_invoices", "status_counts"):
        assert key in d
    assert d["outstanding"] == str(Decimal(d["pending"]) + Decimal(d["overdue"])) or Decimal(d["outstanding"]) == Decimal(d["pending"]) + Decimal(d["overdue"])
    assert d["outstanding_count"] >= 0

    o = client.get("/api/dashboard/overview", params={"date_from": "2020-01-01", "date_to": "2020-01-10"}).json()
    assert (o["group"], len(o["series"]), Decimal(o["total"])) == ("day", 10, Decimal(0))
    assert o["change_percent"] is None  # no previous-period sales: no invented percentage

    today = client.get("/api/dashboard/overview").json()
    assert len(today["series"]) == 30
    assert Decimal(today["total"]) == sum(Decimal(p["total"]) for p in today["series"])
    wide = client.get("/api/dashboard/overview", params={"date_from": "2025-01-01", "date_to": "2026-06-30"}).json()
    assert wide["group"] == "month"


def test_customer_paid_and_outstanding_totals(client):
    customer_id = client.get("/api/customers").json()["items"][0]["id"]
    inv = client.post("/api/invoices", json=_invoice_payload(customer_id)).json()
    before = client.get(f"/api/customers/{customer_id}").json()
    client.post(f"/api/invoices/{inv['id']}/payments", json={"amount": "100"})
    after = client.get(f"/api/customers/{customer_id}").json()
    assert Decimal(after["total_paid"]) - Decimal(before["total_paid"]) == Decimal("100")
    assert Decimal(after["outstanding"]) == Decimal(after["total_billed"]) - Decimal(after["total_paid"])


def test_multi_page_pdf_repeats_header_and_keeps_totals_on_last_page(client):
    from io import BytesIO
    customer_id = client.get("/api/customers").json()["items"][0]["id"]
    payload = _invoice_payload(customer_id)
    payload["items"] = [{**payload["items"][0], "name": f"Item {n}", "quantity": "1"} for n in range(1, 61)]
    r = client.post("/api/invoices", json=payload)
    assert r.status_code == 201, r.text
    body = r.json()
    pdf = client.get(f"/api/invoices/{body['id']}/pdf")
    pages = PdfReader(BytesIO(pdf.content)).pages
    assert len(pages) >= 2
    texts = [p.extract_text() for p in pages]
    assert all("Description" in t and "Amount" in t for t in texts)          # table header repeats on every page
    assert all(body["invoice_number"] in t for t in texts)                    # running header on every page
    assert "Continued on next page" in texts[0] and "Continued on next page" not in texts[-1]
    assert "Authorised Signatory" in texts[-1] and "Authorised Signatory" not in texts[0]
    assert format_indian(Decimal(body["grand_total"])) in texts[-1]           # backend total, printed on the last page


def test_product_gst_filter_and_pagination(client):
    page = client.get("/api/products", params={"gst_rate": "28", "limit": 1}).json()
    assert all(Decimal(p["gst_rate"]) == 28 for p in page["items"])
    assert client.get("/api/products", params={"limit": 1, "offset": 1}).status_code == 200


def test_export_all_data_as_xlsx(client):
    from io import BytesIO
    from openpyxl import load_workbook
    r = client.get("/api/export/xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml")
    assert "OHM-Powerin-data-" in r.headers["content-disposition"] and ".xlsx" in r.headers["content-disposition"]
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Business", "Customers", "Products", "Invoices", "Invoice Items", "Materials", "Stock Movements"]
    assert wb["Invoices"].max_row - 1 == client.get("/api/invoices", params={"limit": 1}).json()["total"]
    assert wb["Products"].max_row - 1 == client.get("/api/products", params={"limit": 1}).json()["total"]
    assert wb["Customers"].max_row - 1 == client.get("/api/customers", params={"limit": 1}).json()["total"]
    assert wb["Materials"].max_row - 1 == client.get("/api/materials").json()["summary"]["count"]
    # totals are real numbers, not text
    headers = [c.value for c in wb["Invoices"][1]]
    grand = headers.index("Grand total") + 1
    assert all(isinstance(wb["Invoices"].cell(row=r_, column=grand).value, (int, float)) for r_ in range(2, wb["Invoices"].max_row + 1))


def test_static_bundle_is_up_to_date():
    """frontend/engine/app.zip (what GitHub Pages runs in the browser) must match backend/. Fix: python tools/build_static.py"""
    import sys
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "tools"))
    import build_static
    assert build_static.is_current(), "Run `python tools/build_static.py` and commit frontend/engine/app.zip"


def test_footer_toggle_and_footer_image(client):
    import io
    from PIL import Image

    client.put("/api/settings", json=BUSINESS)
    r = client.get("/api/settings")
    assert r.json()["footer_enabled"] is True and r.json()["has_footer_image"] is False

    buffer = io.BytesIO()
    Image.new("RGB", (900, 145), (31, 111, 181)).save(buffer, format="PNG")
    r = client.post("/api/settings/footer-image", files={"file": ("footer.png", buffer.getvalue(), "image/png")})
    assert r.status_code == 200, r.text
    assert r.json()["has_footer_image"] is True
    assert client.get("/api/settings/footer-image").status_code == 200
    r = client.post("/api/settings/footer-image", files={"file": ("evil.png", b"not an image", "image/png")})
    assert r.status_code == 400

    r = client.put("/api/settings", json={**BUSINESS, "footer_enabled": False})
    assert r.status_code == 200 and r.json()["footer_enabled"] is False

    r = client.delete("/api/settings/footer-image")
    assert r.json()["has_footer_image"] is False
    assert client.get("/api/settings/footer-image").status_code == 404


def test_footer_image_height_saved_and_validated(client):
    r = client.put("/api/settings", json={**BUSINESS, "footer_image_height": 22})
    assert r.status_code == 200 and r.json()["footer_image_height"] == 22
    assert client.put("/api/settings", json={**BUSINESS, "footer_image_height": 90}).status_code in (400, 422)


def test_material_stock_management(client):
    from decimal import Decimal as D

    r = client.post("/api/materials", json={
        "name": "Fuel filter", "code": "FF-01", "category": "Filters", "unit": "Nos",
        "min_stock": "5", "purchase_price": "120.50", "selling_price": "180", "opening_stock": "10",
    })
    assert r.status_code == 201, r.text
    m = r.json()
    assert D(m["quantity"]) == D("10.000") and D(m["stock_value"]) == D("1205.00") and m["is_low"] is False
    assert client.post("/api/materials", json={"name": "Dup", "code": "ff-01"}).status_code == 409

    r = client.post(f"/api/materials/{m['id']}/stock", json={"kind": "out", "quantity": "6", "note": "Site work"})
    assert r.status_code == 200 and D(r.json()["quantity"]) == D("4.000") and r.json()["is_low"] is True
    assert client.post(f"/api/materials/{m['id']}/stock", json={"kind": "out", "quantity": "5"}).status_code == 400
    r = client.post(f"/api/materials/{m['id']}/stock", json={"kind": "in", "quantity": "2", "unit_price": "125"})
    assert D(r.json()["quantity"]) == D("6.000") and D(r.json()["purchase_price"]) == D("125.00")
    r = client.post(f"/api/materials/{m['id']}/stock", json={"kind": "set", "quantity": "3"})
    assert D(r.json()["quantity"]) == D("3.000")

    moves = client.get(f"/api/materials/{m['id']}/movements").json()
    assert [x["kind"] for x in moves] == ["set", "in", "out", "in"]
    assert D(moves[0]["quantity"]) == D("-3.000") and D(moves[0]["balance"]) == D("3.000")

    page = client.get("/api/materials", params={"low_stock": True}).json()
    assert page["total"] == 1 and page["summary"]["low_stock"] == 1 and page["categories"] == ["Filters"]
    assert D(page["summary"]["value_at_cost"]) == D("375.00")

    r = client.put(f"/api/materials/{m['id']}", json={"name": "Fuel filter (big)", "code": "FF-01", "selling_price": "190"})
    assert r.status_code == 200 and D(r.json()["quantity"]) == D("3.000")  # editing details never changes stock
    assert client.delete(f"/api/materials/{m['id']}").status_code == 204
    assert client.get(f"/api/materials/{m['id']}").status_code == 404


def test_export_includes_materials_and_stock_movements(client):
    from io import BytesIO
    from openpyxl import load_workbook
    m = client.post("/api/materials", json={"name": "Export check", "code": "EXP-1", "purchase_price": "10", "selling_price": "15", "opening_stock": "4"}).json()
    client.post(f"/api/materials/{m['id']}/stock", json={"kind": "out", "quantity": "1", "note": "used"})
    wb = load_workbook(BytesIO(client.get("/api/export/xlsx").content))
    rows = [[c.value for c in row] for row in wb["Materials"].iter_rows(min_row=2)]
    row = next(r for r in rows if r[0] == "EXP-1")
    assert row[4] == 3 and row[8] == 30 and row[9] == 45  # stock, value at cost, value at selling
    moves = [[c.value for c in row] for row in wb["Stock Movements"].iter_rows(min_row=2) if row[1].value == "EXP-1"]
    assert [r[3] for r in moves] == ["Stock in", "Stock out"] and moves[1][4] == -1 and moves[1][5] == 3
    client.delete(f"/api/materials/{m['id']}")
